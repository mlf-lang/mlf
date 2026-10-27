from pathlib import Path
import json, re, statistics, hashlib, collections
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parent; W=R.parent
def read(p): return json.loads((W/p).read_text())
F=read('frozen.json'); B=read('evidence/baseline-final-summary.json')
P=read('evidence/mlf-forced-phase-wall.json'); L=read('profile-late-analysis/summary.json')
G=read('evidence/ghc-eventlog-phase-wall.json'); GRT=read('evidence/ghc-eventlog-rts-observations.json')
labels=P['by_label']; median=B['mlf_ordinary']['median_wall_seconds']; target=B['target_seconds']
D=Document(); S=D.sections[0]
S.page_width=Inches(8.5); S.page_height=Inches(11)
S.top_margin=S.bottom_margin=Inches(.66); S.left_margin=S.right_margin=Inches(.7)
S.header_distance=S.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 st.font.size=Pt(11)
D.styles['Normal'].paragraph_format.line_spacing=Pt(16)
D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Title'].font.size=Pt(25); D.styles['Title'].font.bold=True
D.styles['Title'].paragraph_format.space_after=Pt(12)
D.styles['Subtitle'].font.size=Pt(12); D.styles['Subtitle'].font.italic=False
for st in D.styles:
 for e in list(st._element.xpath('.//w:pBdr')): e.getparent().remove(e)
for name,size in [('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
 D.styles[name].font.size=Pt(size); D.styles[name].font.bold=True
 D.styles[name].paragraph_format.line_spacing=1.08
 D.styles[name].paragraph_format.space_before=Pt(10); D.styles[name].paragraph_format.space_after=Pt(7)
footer=S.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run(); run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)
D.core_properties.title='mlf crates 14 模块编译性能诊断与优化前基线'
D.core_properties.subject='Same-source frozen14 compile profile before optimization'
D.core_properties.author=''; D.core_properties.last_modified_by=''
plain=[]
def p(t='',style=None):
 after_table=len(D._element.body)>1 and D._element.body[-2].tag==qn('w:tbl')
 para=D.add_paragraph(t,style)
 if after_table:para.paragraph_format.space_before=Pt(7)
 plain.append(t); return para
def h(t,level=1): p(t,f'Heading {level}')
def page(t):
 para=p(t,'Heading 1'); para.paragraph_format.page_break_before=True
def note(t):
 x=p(t)
 for r in x.runs:r.font.size=Pt(9)
 return x
def table(headers,rows,widths=None,size=9.5):
 t=D.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 if widths is None: widths=[7.1/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,text in enumerate(headers):t.rows[0].cells[i].text=str(text)
 for row in rows:
  cells=t.add_row().cells
  for i,text in enumerate(row):cells[i].text=str(text)
 for ridx,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ridx==0: rep=OxmlElement('w:tblHeader');trPr.append(rep)
  for i,cell in enumerate(row.cells):
   cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
   for tag,value in [('top','45'),('bottom','45'),('left','85'),('right','85')]:
    e=OxmlElement('w:'+tag);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
   cp.append(margins)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   cp.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if ridx==0 else ('F5F7F9' if ridx%2==0 else 'FFFFFF'));cp.append(shade)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=Pt(13.5)
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in pp.runs:
     r.font.name='Calibri';r.font.size=Pt(size);r.font.bold=(ridx==0);r.font.color.rgb=RGBColor(0,0,0)
     r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 plain.append(' | '.join(headers));plain.extend(' | '.join(map(str,row)) for row in rows)
 return t
def sec(ns,n=6): return f'{ns/1e9:.{n}f}'
def phase_rows(names):
 return [[x,str(labels[x]['count']),sec(labels[x]['exclusive_wall_ns']),sec(labels[x]['inclusive_wall_ns'])] for x in names]

p('mlf crates 14 模块编译性能诊断', 'Title')
p('优化前基线与第一项优化决策', 'Subtitle')
p('2026 年 10 月 26 日  |  PR25 aa11ffd 与已批准本地修复')
h('结论')
p(f'同一批 14 个源模块、依赖已编译并复用、目标推断缓存全冷时，普通 MLF 编译中位数为 {median:.9f} s；GHC 9.14.1 plain -O2 的中位数为 9.047647972 s。当前比例为 {B["primary_ratio"]:.6f} 倍，尚未达到不超过 3 倍的目标。达标阈值是 {target:.9f} s，意味着还需约 {median/target:.2f} 倍加速，减少约 {(1-target/median)*100:.2f}% 的普通 wall time。')
table(['主要指标','结果','解释'],[
 ['普通 MLF 中位数','163.896307032 s','3 次独立进程'],
 ['GHC plain -O2 中位数','9.047647972 s','同一 14 模块'],
 ['MLF 与 GHC 比例','18.114797 倍','目标不超过 3 倍'],
 ['每次目标推断','544 fresh / 0 reused','不是缓存命中跑分'],
], [2.7,1.9,2.5],10)
h('第一项优化应从命名空间成员选择开始',2)
p('late cost-centre profile 中，Bootstrap.Module 的 partialSelectItem 与 selectType 在四个互不重叠的 CCS 上下文中占 38.4513% sampled CPU、44.0067% profiled allocation。直接归属于这些选择器的 Map 分配为 127.280 GB，占全部 Map-exclusive 分配的 71.4565%。主要 Map 热点不能简单归因于 solver。')
p('建议先测 parent filter before union：先按父实体筛选成员，再 union 与去重，避免每个成员选择都构造完整 visible interface。必须保持实体排序、alias 去重、成员可见性、歧义与缺失名称错误、循环导出 fixed point、Abstract 成员惰性和全部独立检查。此处是有证据的候选，不是已验证收益。')
p('本报告在任何性能优化代码修改之前形成。forced phase wall 与 late CPU profile 用于定位，不用来替代普通基线。GHC 输出优化后的 native objects，而 MLF 输出经独立检查的 xMLF crate 并写入 query cache；两者完成的产品不同。')
note('阅读顺序：第 1 至 2 节界定基线与诊断扰动；第 3 至 7 节解释热点和验证门槛；附录给出阶段数值和复现证据。GB 为十进制 10^9 bytes，MiB 为 2^20 bytes。')

page('1 比较范围与普通基线')
p('计时从新子进程启动到退出。准备依赖、复制输入、哈希校验和 benchmark harness 的前后检查不计入编译 wall time；编译进程内必要的解析、推断、检查与目标输出计入。没有执行完整 74 模块 native bootstrap，本报告也不作该范围的性能或语义结论。')
table(['配置','第 1 次 s','第 2 次 s','第 3 次 s','中位数 s'],[
 ['MLF 普通',*['%.9f'%x for x in B['mlf_ordinary']['wall_seconds']],f'{median:.9f}'],
 ['GHC plain -O2',*['%.9f'%x for x in B['ghc_plain_O2']['wall_seconds']],'9.047647972'],
 ['GHC repo flags',*['%.9f'%x for x in B['ghc_repo_demand_flags']['wall_seconds']],'8.847117296'],
], [1.65,1.36,1.36,1.36,1.37],9)
p('主比较必须使用 plain -O2。repo flags 额外包含 -fpedantic-bottoms -fno-state-hack，仅作为独立敏感性对照，其 8.847117296 s 不能冒充 plain 基线。早期 baseline-summary.json 记录的是该对照；最终主口径在 baseline-final-summary.json 与三份 baseline-plain-ghc receipt 中。')
h('控制条件',2)
p('GHC 命令核心为 --make -no-link -j1 -O2 -XHaskell2010，固定 package db、显式 package 和 source roots；plain 依赖也用匹配的 plain flags 预编译后复用。MLF producer 按仓库的 -O2 与 demand-preserving flags 构建，运行 --compile-crate，使用全新的目标 queries 目录，复用冻结的依赖 crate 与 Prelude。')
p('共同环境固定 GHC_ENVIRONMENT=-、LC_ALL=C.UTF-8、OMP_NUM_THREADS=1、OPENBLAS_NUM_THREADS=1、MAKEFLAGS=-j1；普通样本 GHCRTS 为空。任务控制的重负载串行运行，但没有 CPU affinity 或独占物理宿主保证。OS cache 未清空；这里的 cold 指目标推断结果全冷，不是冷页缓存。')
p('宿主快照为 Linux 6.18.55 x86_64、AMD EPYC 9V74，约 9.7 GiB 物理内存、无 swap。cgroup 控制未挂载，额外的宿主资源 cap 未知。资源守护上限为单进程 6 GiB address space、进程树 5 GiB RSS；宿主可用内存下限 3 GiB、磁盘余量下限 4 GiB、输出增长上限 8 GiB、wall 上限 6 小时。')
h('统计范围',2)
p('每组 3 个样本只描述当前宿主。MLF 范围为 163.331291426–170.168008194 s，MAD 为 0.565015606 s；plain GHC 范围为 8.453523167–9.532567807 s。没有据此建立置信区间或跨机器推广。历史 128.58 s / 5.67 s 的命令与 revision 未独立核验，未合并入基线。')
note('第 1 次 MLF 普通样本早于 CPU/rusage 仪表加入，CPU 数据缺失，未回填估算；其输入、544/0 计数和产物字节一致性已独立后验核对。')

page('2 测量口径与诊断扰动')
table(['测量','进程 wall s','相对普通 MLF','允许的解释'],[
 ['普通 MLF 中位数','163.896307032','1.000000','是否达标'],
 ['普通实现加 RTS stats','163.412719949','0.997049','整体 GC 与 allocation'],
 ['selected-output forced','197.156816802','1.202936','诊断性阶段 wall'],
 ['late cost-centre','349.293843396','2.131188','sampled CPU 与 allocation'],
],[2.1,1.45,1.35,2.2],9.5)
p('普通 RTS 版本保留 qualified producer 的 250 个非 Main .hi/.o 文件且逐字节相同，只重建 Main wrapper 并重新链接启用 RTS options。单次 wall 比普通中位数低约 0.30%，落在观察到的噪声尺度内，不能宣称统计功能带来负开销。late profile 与 forced profile 则明显改变执行成本，不能用一个总倍率校正各阶段。')
h('普通与 forced 的 RTS 对照',2)
table(['RTS 指标','普通实现','forced 诊断'],[
 ['allocated bytes','432,025,763,688','455,360,043,208'],
 ['copied during GC bytes','15,427,911,136','15,532,515,824'],
 ['max residency bytes','61,927,160','1,738,370,216'],
 ['major GC residency 样本数','146','43'],
 ['GC elapsed s','17.048','20.801'],
 ['MUT elapsed s','146.353','176.234'],
 ['RTS total elapsed s','163.401','197.036'],
 ['采样进程树峰值 RSS bytes','245,940,224','4,074,078,208'],
],[3.0,2.05,2.05],9.5)
p('RTS 时间按输出精度取整，process wall 则来自纳秒单调时钟；两者边界不同，不强行相等。普通 RTS 的 GC CPU 为 17.054 s，MUT CPU 为 141.144 s，total CPU 为 158.200 s。该进程的 gen0 / gen1 收集次数为 101,071 / 146。')
p('forced forcing 增大存活数据与内存占用，改变惰性求值、保留和 GC。forced 最大 residency 约 1.738 GB、进程树 RSS 约 4.074 GB；普通进程树 RSS 约 235 MiB。allocation 是累计分配，不是常驻内存；max residency 是 GC 观察，RSS 是另一种指标。')
note('GC 是 wall time 内的正交 overlay，不能再作为阶段加总。尤其不能从 forced 阶段减去 forced-again 就宣称恢复了普通阶段时间。')

page('3 Forced 阶段显示工作落在哪里')
p('本节全部秒数属于 selected-output forced diagnostic，不是普通编译阶段 wall。24,057 个区间完成且 root 状态为 ok；每个 exclusive 等于本体区间减去直接子区间，本体边界和子 helper 开销留在父级 exclusive。')
top=[k for k in labels if '/forced-again' not in k][:17]
table(['阶段标签','次数','exclusive s','inclusive s'],phase_rows(top),[3.5,.6,1.5,1.5],9)
p('全体 exclusive 精确合计 195.464745170 s，与 diagnostic-total root 完全一致。进程总 wall 197.156816802 s = root 195.464745170 s + root 外残差 1.692071632 s；残差覆盖初始化、未纳入 root 的工作和缓冲输出等，不能全部命名为某一个阶段。')
p('solver.final 的 inclusive 为 42.628325615 s；inference.fragment 的 inclusive 为 70.993421047 s。它们已包含子工作，因此不能把这些 inclusive 与表中的子阶段再次相加。按标签累计次数也不等于源码函数调用次数。')
note('完整非重复标签按流水线列在附录；完整逐事件记录保留在 runs/profile-forced-1/phases.jsonl，不在正文复制数万行。')

page('4 复合阶段不能按名字直接下结论')
h('module resolve 包含命名空间与导出工作',2)
p('module.resolve 的 51.682578694 s 是强制 resolved crate 结果的复合区间，覆盖 module resolution、导入导出环境与 fixed-point 导出处理等被该结果 demand 的工作。late CCS 独立证实大量 selector 成本落在 settleExports 和 exportsFor 路径；两种证据相互解释，但不能把 sampled CPU shares 换算为这里的 wall 秒数。')
h('cache write publish 不只是文件写入',2)
p('cache.write-publish 的 exclusive 为 23.019529905 s，inclusive 为 28.157811578 s。嵌套 query.cache.encode 的 5.137713778 s 已单独记账并从父 exclusive 扣除。父级还包含 writeFile、参数与 UTF-8 处理，以及外部 /bin/mv 的 fork、exec、wait。runtime/mlf_system.h 的 os_runProcess 使用 10 ms polling；不是 supervisor 的 5 ms loop。')
p('forced 版本的大 heap 可能放大 fork 成本，所以这不是“普通缓存写入要 23 秒”的证据。既不能把该标签全部归为磁盘吞吐，也不能把 fixed polling interval 乘以 544 当作已经测出的实际等待总量。需要后续在普通实现中拆分发布路径，才能判断优化空间。')
h('crate assemble validate 不等于独立 checker',2)
p('crate.assemble-validate 的 exclusive 为 12.222363884 s，inclusive 为 17.667390839 s。它包括最终 resolvedExports 再次执行 exportsFor、exports 去重、依赖 semantic hash、layout / reachability / bindings / protection、makeObject 与 validateInterface，以及所选 CrateObject 和 QueryReport 的 deep force。延迟的 GroupResult semantic hash 也可在这里被 demand。')
p('嵌套 payload fingerprint 与 encode 的 5.445025443 s 已从此父 exclusive 扣除。把剩余 12.222 s 叫作“类型检查器成本”会误导优化。late profile 中可识别的独立 checker 根只有约 1.55% sampled CPU；它不是当前首要目标，也不授权减弱检查。')
h('forcing 边界已检查 但不承诺普通求值顺序',2)
p('PhaseSelfTest、ForceSelfTest 与区间解析控制均通过；optimized Core 中 opaque purePhase / purePhaseOnce 的 value 参数在 helper 边界保持 lazy，forceAction 在 IO 中执行所选 forcing。该证据检查 helper 优化行为，不是对所有 callsite 普通 demand 的普遍证明。强制指定输出本来就会改变 demand 和 retention。')
p('具体 demand 边界也已逐处复核：Constraint 的 incremental 输入 graph 在 prepare 外强制，费用归 constraints.generate；SourcePlan 只 force 选定输出，不触碰未使用的 legacy lowering；TypeContext memo 的 demand 仍留在 evidence / elaboration；String forcing 不强制 lazy index；solver 跳过未使用的 ScopeIndex jump memo。这些是诊断边界约束，不是普通求值顺序的等价证明。')
note('原始 RTS before/after snapshot 可用于审计，但 GC 更新的计数器不能当成精确 per-phase allocation 或 phase-end retained heap。')

page('5 Late CPU profile 指向的实际热点')
p('完整解析 81,682 个最终 CCS 行；exclusive 整数总和精确等于 318,156 ticks 与 431,827,326,224 allocated bytes。实际 header 为 1,000 µs/tick；-i0.005 是 heap profiling interval 参数，不是 time-profiler tick 请求，不能据此换算 CPU 样本。20 个 flat exact 行和全部 inherited 百分比在输出舍入范围内对得上。')
regs=L['regions_not_additive']
items=[('成员选择器并集',L['selection_union_disjoint']),('final solver',regs['final_solver']),('fingerprint 根',regs['fingerprint']),('依赖 closure 加载',regs['dependency_interface_closure']),('独立 checker 可见根',regs['independent_checker_surviving_checkClosedWithBindings'])]
table(['区域','exact ticks','CPU share','allocation GB','alloc share'],[[n,f'{v["ticks"]:,}',f'{v["tick_share"]*100:.4f}%',f'{v["bytes"]/1e9:.3f}',f'{v["allocation_share"]*100:.4f}%'] for n,v in items],[2.35,1.1,1.15,1.3,1.2],9.5)
p('这些是解释热点的不同区域视图，不构成可随意相加的总账。特别是 module_resolution 与 selector 子树、solver 与 instantiation 子树、Map 与 balance 子树可能重叠。报告只在已证明互斥的集合内加总。122.335 sampled CPU seconds 只是 122,335 ticks × 1 ms，既不是普通 wall，也不是 exact phase wall。')
h('四个 selector 上下文并不代表四次运行调用',2)
p('partialSelectItem 的两个 settleExports 上下文共 75,527 ticks、118.029 GB；selectType 在 resolution 和 final resolvedExports 两处共 46,808 ticks、72.004 GB。后期插入与 inlining 会改变可见 CCS，entry counts 不能直接作为调用数或 fixed-point 迭代次数。sourceMethodVisible 没有 surviving named CCS，不等于它的成本为零。')
h('final solver 的无重叠分区',2)
table(['solver 分区','ticks','allocation GB'],[[x['label'],f'{x["ticks"]:,}',f'{x["bytes"]/1e9:.3f}'] for x in L['final_solver_disjoint_partition']],[4.2,1.3,1.6],9.5)
p('instantiation 已包含其内部 nested unify；top_level_unify 是另一个互斥分区。solver 总计为 76,462 ticks、75.689 GB，占全 profile 24.0329% CPU，而非“所有 Map 成本都属于 solver”。')

page('6 GHC 的阶段账本与 GC')
p('GHC eventlog 诊断使用 repo demand flags，不是主分母的 plain -O2 配置。单次进程 wall 为 8.855613232 s，接近该配置的普通中位数 8.847117296 s，但单次差异不能量化 eventlog 的无偏开销。')
agg=collections.defaultdict(lambda:[0,0,0])
for k,v in G['by_label'].items():
 name=re.sub(r' \[[^]]+\]$','',k)
 for i,field in enumerate(['count','exclusive_wall_ns','inclusive_wall_ns']):agg[name][i]+=v[field]
sorted_agg=sorted(agg.items(),key=lambda x:-x[1][1]); topghc=sorted_agg[:17]
grow=[]
for name,v in topghc:
 short=name
 if name.startswith('Float out('):short='Float out  JoinsToTop=True'
 grow.append([short,f'{v[0]:,}',sec(v[1]),sec(v[2])])
rest=[sum(v[i] for k,v in sorted_agg[17:]) for i in range(3)]
grow.append(['其余已标记标签',f'{rest[0]:,}',sec(rest[1]),'不合并父子 inclusive'])
table(['阶段族 合并模块后缀','区间数','exclusive s','inclusive s'],grow,[3.35,.7,1.35,1.7],9)
p('39,890 个区间的 exclusive 总和为 8.496111323 s，已标记区间覆盖之外的残差为 0.359501909 s；合计精确还原 8.855613232 s 进程 wall。CodeGen inclusive 包含后端子阶段；NCG 与 Cmm 不能重复加到它上面。')
p('从 201,362 个事件中选取 USER_MSG / traceEventIO 的 tag 19 相位通道，79,780 markers 配对为 39,890 phases；另有 tag 58 diagnostic messages 1,012 条，不与相位通道混用。parser 对 block 长度、最终 block、双通道和区间嵌套有控制测试。')
note('GC 525 个区间的 elapsed overlay 为 2.176201864 s，已包含在阶段和进程 wall 中。HEAP_ALLOCATED 最后观测为 10,593,411,424 bytes；HEAP_LIVE 最大观测为 63,053,704 bytes，仅是 10 次 major-GC live 观测，不是连续峰值或 RSS。不能从这些时间戳推断 GC CPU、精确 mutator time 或各阶段 allocation。')

page('7 验证结论与下一步门槛')
h('已经建立的事实',2)
p('输入 provenance audit 核验了 2,581 个唯一 pinned 文件，涵盖真实 GHC binary、launcher、settings、toolchain、installed dependencies、resolved source modules、host/configuration inputs、14 个冻结源模块、GHC objects 与 MLF dependency artifacts。profile 完成后的再次审计也全部通过。哈希校验约 1.259 GB 输入；这证明输入来源一致，不是独立的完整语义测试。')
p('普通 MLF 三次、RTS、forced 与 late-profile 目标编译均成功，544 fresh / 0 reused；.mlfci 和 .mlfco 与先前参考输出逐字节一致。GHC 普通样本确认恰好编译 14 模块且依赖对象没有改变。此处没有绕过独立 xMLF 检查。')
h('第一项候选的验证门槛',2)
for t in [
 '先保留并发布本报告、冻结输入与原始数据，再改变性能代码；不把诊断 instrumentation 当作优化。',
 '只实施 parent filter before union 的最小改动；保持 entity identity、排序、alias 去重、错误与 lazy Abstract 语义。优先利用现有 selector 语义，不更换 hash 或跳过 checks。',
 '覆盖显式成员列表、全部成员、Abstract、重复 alias、歧义/缺失名称、成员可见性、导出循环与 fixed point 的正反例；保留模块与 crate 的独立验证。',
 '同一 frozen14 冷目标重编译必须仍为 544/0，并比较完整产物字节、query semantic fingerprint 与错误行为；随后运行授权范围内更广的回归测试。',
 '在没有 profiler 的普通 executable 上重新独立计时，至少使用与基线相同的样本数、环境和资源守护；依据普通 median 验收收益，同时报告 allocation / RSS 是否回退。',
 '只有普通 MLF median ≤ 27.142943916 s 才在本宿主本输入上达到目标。一次局部改动不能预先保证整体 6.04 倍加速。',
]: p('• '+t)
h('后续候选与不做的事情',2)
p('第二优先级是 solver instantiation 中重复 index/map 构建与 projection/restoration；第三是不可变 serialization / fingerprint 共享；第四是在一次 compile 内复用已认证依赖快照。20 个 fingerprint 根合计 36,344 ticks、55.591 GB；两处依赖加载分别为 3,725 / 3,756 ticks、各约 6.262 GB。都需要代码与数据身份论证及普通计时，不采用弱 hash、过期缓存或减少检查来换速度。')

page('附录 A Forced 输入与推断阶段')
p('以下是按工作类别排列的同一次 forced 诊断。exclusive 可相加；inclusive 仅用于看父区间覆盖范围，不能与子区间相加。表中保留 6 位小数，精确纳秒见 JSON。')
src=['dependency.initial.load-decode-validate','dependency.prelude.load-decode-validate','source.load-all','source.parse','source.clean','crate.prepare-selected-outputs','module.resolve','source.plan.prepare','source.group.prepare','source.group.finalize']
table(['标签','次数','exclusive s','inclusive s'],phase_rows(src),[3.6,.5,1.5,1.5],9)
inf=['inference.fragment','constraints.generate','solver.incremental.prepare','solver.incremental.unify','solver.incremental.steps','solver.incremental.finish-audit','solver.final','solver.final.prepare','solver.final.unify','solver.final.steps','solver.final.finish-normalize-audit','elaboration','check.compiler-core','evidence.resolve','fragment.parameter-lowering','check.fragment-payload','agreement.type','agreement.erasure','fragment.collect-bindings']
table(['标签','次数','exclusive s','inclusive s'],phase_rows(inf),[3.6,.5,1.5,1.5],9)

page('附录 B Forced 缓存与产物阶段')
other=[k for k in labels if '/forced-again' not in k and k not in src+inf]
table(['标签','次数','exclusive s','inclusive s'],phase_rows(other),[3.6,.5,1.5,1.5],9)
repeats=[(k,v) for k,v in labels.items() if '/forced-again' in k]
repeat_ns=sum(v['exclusive_wall_ns'] for k,v in repeats)
repeat_count=sum(v['count'] for k,v in repeats)
p(f'全部 forced-again 标签合计 {repeat_count:,} 个区间、{repeat_ns/1e9:.9f} s exclusive。它们刻画再次遍历已 force 结果的诊断工作，不是完整 instrumentation 开销，更不是可以直接从普通编译中扣除的“冗余”。')
sum_non=sum(v['exclusive_wall_ns'] for k,v in labels.items() if '/forced-again' not in k)
table(['精确总账','纳秒','秒'],[
 ['非 forced-again exclusive',f'{sum_non:,}',f'{sum_non/1e9:.9f}'],
 ['forced-again exclusive',f'{repeat_ns:,}',f'{repeat_ns/1e9:.9f}'],
 ['root 全体 exclusive',f'{P["root_wall_ns"]:,}',f'{P["root_wall_ns"]/1e9:.9f}'],
 ['root 外 process residual','1,692,071,632','1.692071632'],
 ['process wall','197,156,816,802','197.156816802'],
],[3.25,2.0,1.85],9.5)
p('query.source-key.encode、query.key.digest、query.cache.encode、artifact.payload.encode 与 fingerprint 是不同 demand 边界。编码和 hash 可以间接 demand 尚未求值的数据；因此名称表示计时包围的动作及其所需工作，不保证是纯字符串转换或纯 hash CPU。')

page('附录 C 输入身份与可复现路径')
p('所有下述相对路径以 mlf-crates-profile-work-20261026 为根目录 W。复现必须使用原 receipt 的完整 argv、cwd、environment、compiler hash 与匹配依赖，不能只复制本报告中的缩写命令。源码与诊断 patch、原始日志和 JSON 应随报告一起保留。')
for name,val in [('Base commit',F['base_commit']),('Approved input patch SHA256',F['input_patch_sha256']),('Frozen manifest SHA256',hashlib.sha256((W/'frozen.json').read_bytes()).hexdigest()),('Baseline MLF compiler SHA256',F['compiler_sha256']),('GHC launcher SHA256',F['ghc_sha256']),('.mlfci reference SHA256',F['reference_artifacts']['.mlfci']),('.mlfco reference SHA256',F['reference_artifacts']['.mlfco'])]:
 x=p(name);x.runs[0].bold=True;x.runs[0].font.size=Pt(10)
 x.paragraph_format.space_after=Pt(2);x.paragraph_format.line_spacing=Pt(13)
 x=note(val);x.paragraph_format.line_spacing=Pt(13);x.paragraph_format.space_after=Pt(5)
h('固定的 14 个目标模块',2)
mods=F['order']
table(['编译顺序 1 至 7','编译顺序 8 至 14'],[[mods[i],mods[i+7]] for i in range(7)],[3.55,3.55],9.5)
h('复现命令的入口',2)
p('普通 MLF：recipes/benchmark.py mlf <新标签>。plain GHC：recipes/benchmark.py ghc <新标签> --plain-o2。两者要求目录尚不存在，并且执行前的资源 guard 应与原运行一致。不要在保留结果的既有标签中重跑，不要在计时期间并发构建或渲染文档。')
p('阶段解析：diagnostic/analyze_phase.py；GHC eventlog 解析：recipes/ghc_eventlog_phases.py 与 ghc_eventlog_rts.py；late CCS 重建：profile-late-analysis/analyze_late_profile.py。逐事件和树数据用程序处理，而非复制到文档。')

page('附录 D 证据索引与已知限制')
sources=[
 ('主基线和 RTS 对照','evidence/baseline-final-summary.json'),
 ('普通样本完整命令与结果','runs/baseline-mlf-{1,2,3}/receipt.json\nruns/baseline-plain-ghc-{1,2,3}/receipt.json'),
 ('repo flags 敏感性样本','evidence/baseline-summary.json\nruns/baseline-make-ghc-{1,2,3}/receipt.json'),
 ('普通实现 RTS','runs/profile-baseline-rts-1/receipt.json 与 stderr'),
 ('forced 完整事件与总账','runs/profile-forced-1/phases.jsonl\nevidence/mlf-forced-phase-wall.json'),
 ('late CCS exact counts 与原树','profile-late-analysis/summary.json\nprofile-late-analysis/ccs-tree.json.gz'),
 ('GHC eventlog 阶段与 RTS','evidence/ghc-eventlog-phase-wall.json\nevidence/ghc-eventlog-rts-observations.json'),
 ('输入 pins 与 provenance','frozen.json\nevidence/provenance-audit-before-profile.json\nevidence/provenance-audit-after-profile.json'),
 ('forcing 与 helper 验证','diagnostic-tests/receipt.json\nevidence/phase-helper-optimized-core.json'),
 ('复合阶段 callsite 审计','evidence/mlf-forced-composite-interpretation.json'),
 ('资源保护与取消记录','guards/*/bounded-receipt.json\nevidence/auto-profile-cancel-observation.json'),
]
table(['证据用途','相对路径'],sources,[2.2,4.9],9)
h('限制需与结果一起保留',2)
for t in [
 '三次普通样本、单次诊断不能给出普遍分布或跨版本/机器结论；普通与诊断数据不混样，不以 CPU 百分比乘 ordinary wall。',
 '早期 automatic SCC profile 在约 321.9 s 被取消，未产生代表性完成结果，排除于计时基线。终端退出码为 130，后验观察输出未继续变化；可见性检查仅限当前进程 namespace。',
 '最早一批 guard 的 output-growth root 只覆盖日志，未覆盖实际运行输出；wall、RSS、address space、宿主内存与磁盘余量保护仍有效，后续已修正。未计时的 plain dependency setup 重复 output-root 参数只选择了最后一个 root；实际 setup 输出 16,873,890 bytes 事后核查。',
 'checked crate 字节一致和诊断 controls 提供本任务范围内证据，不替代全项目测试或完整 native bootstrap。报告没有测得任何优化收益，也没有修改算法、减弱验证或改变编译语义。',
]: p('• '+t)

out=R/'mlf-crates14-preoptimization-profile-zh.docx';D.save(out)
(R/'mlf-crates14-preoptimization-profile-zh.txt').write_text('\n\n'.join(plain)+'\n')
facts={'primary_ratio':B['primary_ratio'],'target_seconds':target,'forced_root_ns':P['root_wall_ns'],'forced_nonrepeat_exclusive_ns':sum_non,'forced_repeat_exclusive_ns':repeat_ns,'forced_repeat_count':repeat_count,'ghc_aggregated_exclusive_ns':sum(v[1] for v in agg.values()),'ghc_phase_count':G['phase_count'],'report_sources':{str(x):hashlib.sha256((W/x).read_bytes()).hexdigest() for x in ['frozen.json','evidence/baseline-final-summary.json','evidence/mlf-forced-phase-wall.json','evidence/ghc-eventlog-phase-wall.json','evidence/ghc-eventlog-rts-observations.json','profile-late-analysis/summary.json','evidence/provenance-audit-after-profile.json','evidence/mlf-forced-composite-interpretation.json']}}
(R/'report-derived-facts.json').write_text(json.dumps(facts,indent=2)+'\n')
(R/'README-engineering.txt').write_text('mlf-crates14 优化前工程摘要\n\n普通 MLF median 163.896307032 s；GHC 9.14.1 plain -O2 median 9.047647972 s；ratio 18.114797；目标 <=27.142943916 s。\n同一冻结14源，依赖复用，目标544 fresh / 0 reused。GHC native objects 与 MLF checked xMLF+query cache 的终点不对称。\n\n首项候选：Bootstrap.Module partialSelectItem/selectType 的 parent filter before union。late CCS 四个不重叠上下文占38.4513% sampled CPU、44.0067% allocation。不是已验证收益；保留alias、排序、去重、错误、fixed point、惰性与所有checks。\n\n禁止混淆：forced diagnostic 197.156816802 s 不等于 ordinary phase wall；late profile 349.293843396 s / 318156 ticks 不可按CPU占比乘普通wall；GC是overlay，不加到阶段之上。主GHC分母是plain -O2，不是repo extra flags的8.847117296 s。\n\n详细可编辑报告：mlf-crates14-preoptimization-profile-zh.docx\n同内容纯文本：mlf-crates14-preoptimization-profile-zh.txt\n复算值与来源哈希：report-derived-facts.json\n报告构建入口：build_report.py（只读原证据并写本目录，不运行编译）\n\n优化前状态：无性能代码改动。本报告须先交付，再按相同普通计时与语义门槛验证最小候选。未执行完整74模块native bootstrap。\n')
print(out)
